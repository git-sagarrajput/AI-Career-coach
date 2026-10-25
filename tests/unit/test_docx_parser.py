from __future__ import annotations

import io

import docx
from _doc_builders import add_bookmark_link, add_hyperlink

from career_coach.ingestion.docx_parser import parse_docx


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_t10_order_kinds_and_heading_styles():
    d = docx.Document()
    d.add_heading("Experience", level=1)
    d.add_paragraph("Engineer at Example Corp")
    d.add_paragraph("Built APIs in Python", style="List Bullet")
    d.add_paragraph("")
    raw = parse_docx(_save(d))
    assert [(b.kind, b.text) for b in raw.blocks] == [
        ("heading", "Experience"),
        ("paragraph", "Engineer at Example Corp"),
        ("list_item", "Built APIs in Python"),
    ]
    assert raw.blocks[0].style_name == "Heading 1"


def test_t11_table_with_merged_cell():
    d = docx.Document()
    table = d.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"r{r}c{c}"
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "merged"
    raw = parse_docx(_save(d))
    rows = [b for b in raw.blocks if b.kind == "table_row"]
    assert len(rows) == 3
    assert rows[0].text == "merged | r0c2"
    assert raw.table_count == 1


def test_t12_hyperlinks_and_bookmark():
    d = docx.Document()
    add_hyperlink(d.add_paragraph("Profile: "), "https://github.com/demo", "GitHub")
    add_bookmark_link(d.add_paragraph(), "section2", "Jump")
    table = d.add_table(rows=1, cols=1)
    add_hyperlink(table.cell(0, 0).paragraphs[0], "https://portfolio.example.com", "Portfolio")
    raw = parse_docx(_save(d))
    urls = {link.url_raw: link for link in raw.links}
    assert set(urls) == {"https://github.com/demo", "https://portfolio.example.com"}
    assert urls["https://github.com/demo"].sources[0].anchor_text == "GitHub"
    assert urls["https://github.com/demo"].sources[0].kind == "docx_relationship"
