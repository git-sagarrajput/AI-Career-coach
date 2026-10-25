from __future__ import annotations

import io
import logging

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from career_coach.models.documents import (
    ExtractedLink,
    LinkSource,
    ProcessingWarning,
    RawBlock,
    RawExtraction,
)
from career_coach.models.enums import Severity

logger = logging.getLogger(__name__)
_COMPONENT = "docx_parser"


def _paragraph_kind(paragraph: Paragraph) -> str:
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style_name.startswith("Heading") or style_name == "Title":
        return "heading"
    p_pr = paragraph._p.pPr
    has_numbering = p_pr is not None and p_pr.find(qn("w:numPr")) is not None
    if "List" in style_name or has_numbering:
        return "list_item"
    return "paragraph"


def _iter_body(document: DocxDocument):
    if hasattr(document, "iter_inner_content"):
        yield from document.iter_inner_content()
        return
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _row_texts(table: Table, warnings: list[ProcessingWarning], depth: int = 0) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        seen_tc: set[int] = set()
        cells: list[str] = []
        nested_rows: list[str] = []
        for cell in row.cells:
            if id(cell._tc) in seen_tc:
                continue
            seen_tc.add(id(cell._tc))
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            if text:
                cells.append(text)
            for nested in cell.tables:
                if depth == 0:
                    nested_rows.extend(_row_texts(nested, warnings, depth + 1))
                else:
                    warnings.append(
                        ProcessingWarning(
                            code="nested_table_truncated",
                            severity=Severity.INFO,
                            component=_COMPONENT,
                            message="A deeply nested table was skipped.",
                        )
                    )
        if cells:
            rows.append(" | ".join(cells))
        rows.extend(nested_rows)
    return rows


def _paragraph_links(paragraph: Paragraph) -> list[ExtractedLink]:
    links: list[ExtractedLink] = []
    for hyperlink in getattr(paragraph, "hyperlinks", []):
        address = (hyperlink.address or "").strip()
        if not address:
            continue
        anchor = (hyperlink.text or "").strip()[:200] or None
        links.append(
            ExtractedLink(
                url_raw=address[:2048],
                sources=[LinkSource(kind="docx_relationship", anchor_text=anchor)],
            )
        )
    return links


def parse_docx(data: bytes) -> RawExtraction:
    document = docx.Document(io.BytesIO(data))
    blocks: list[RawBlock] = []
    links: list[ExtractedLink] = []
    warnings: list[ProcessingWarning] = []
    table_count = 0

    header_texts: list[str] = []
    for section in document.sections:
        for part in (section.header, section.footer):
            if part is None or (part.is_linked_to_previous and header_texts):
                continue
            for paragraph in part.paragraphs:
                text = paragraph.text.strip()
                if text and text not in header_texts:
                    header_texts.append(text)
                    links.extend(_paragraph_links(paragraph))
    for text in header_texts:
        blocks.append(RawBlock(text=text, kind="paragraph", contact_hint=True))

    for index, item in enumerate(_iter_body(document)):
        if isinstance(item, Paragraph):
            links.extend(_paragraph_links(item))
            text = item.text.strip()
            if not text:
                continue
            blocks.append(
                RawBlock(
                    text=text,
                    kind=_paragraph_kind(item),
                    paragraph_index=index,
                    style_name=item.style.name if item.style is not None else None,
                )
            )
        elif isinstance(item, Table):
            for row in item.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        links.extend(_paragraph_links(paragraph))
            for row_index, row_text in enumerate(_row_texts(item, warnings)):
                blocks.append(
                    RawBlock(text=row_text, kind="table_row", table_index=table_count, row_index=row_index)
                )
            table_count += 1

    known = {link.url_raw for link in links}
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/hyperlink") and rel.is_external:
            target = (rel.target_ref or "").strip()
            if target and target not in known:
                known.add(target)
                links.append(
                    ExtractedLink(url_raw=target[:2048], sources=[LinkSource(kind="docx_relationship")])
                )

    xml = document.element.xml
    if "w:txbxContent" in xml:
        warnings.append(
            ProcessingWarning(
                code="docx_textbox_content_skipped",
                severity=Severity.WARNING,
                component=_COMPONENT,
                message="Text inside Word text boxes or shapes was not extracted.",
            )
        )
    if len(document.inline_shapes) > 0:
        warnings.append(
            ProcessingWarning(
                code="docx_images_ignored",
                severity=Severity.INFO,
                component=_COMPONENT,
                message="Images in the document were ignored.",
            )
        )
    return RawExtraction(
        blocks=blocks, links=links, page_count=None, table_count=table_count, warnings=warnings
    )
