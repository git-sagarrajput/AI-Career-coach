"""Synthetic PDF/DOCX builders shared by tests and scripts/generate_sample_data.py.

Content format (JSON):
{
  "name": "...", "contact": ["line", ...],
  "links": [{"text": "...", "url": "..."}],
  "sections": [{"heading": "...", "paragraphs": ["..."], "bullets": ["..."], "table": [["a", "b"], ...]}],
  "hidden_text": "optional white text (PDF only)",
  "image_only_page": false
}
Sections with "heading": null render their content without a heading.
"""

from __future__ import annotations

import io
import textwrap
from typing import Any

import docx
import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_LEFT = 56
_BOTTOM = 790
_LINE = 14
_WRAP = 92


# ------------------------------------------------------------------ PDF


class _PdfWriter:
    def __init__(self) -> None:
        self.doc = pymupdf.open()
        self.page = self.doc.new_page()
        self.y = 60.0

    def ensure(self, height: float) -> None:
        if self.y + height > _BOTTOM:
            self.page = self.doc.new_page()
            self.y = 60.0

    def line(
        self, text: str, size: float = 10.5, bold: bool = False, color=(0, 0, 0), indent: float = 0
    ) -> None:
        self.ensure(_LINE)
        font = "hebo" if bold else "helv"
        self.page.insert_text((_LEFT + indent, self.y), text, fontsize=size, fontname=font, color=color)
        self.y += _LINE if size < 13 else _LINE + 6

    def wrapped(self, text: str, bullet: bool = False) -> None:
        lines = textwrap.wrap(text, _WRAP - (4 if bullet else 0), break_on_hyphens=False) or [""]
        for i, part in enumerate(lines):
            prefix = "• " if bullet and i == 0 else ("  " if bullet else "")
            self.line(prefix + part, indent=0 if i == 0 or not bullet else 6)

    def gap(self, size: float = 8) -> None:
        self.y += size

    def table(self, rows: list[list[str]]) -> None:
        col_w = 480 / max(len(r) for r in rows)
        height = 18
        self.ensure(height * len(rows) + 4)
        top = self.y - 12
        for r_index, row in enumerate(rows):
            y0 = top + r_index * height
            for c_index, cell in enumerate(row):
                x0 = _LEFT + c_index * col_w
                rect = pymupdf.Rect(x0, y0, x0 + col_w, y0 + height)
                self.page.draw_rect(rect, color=(0, 0, 0), width=0.6)
                self.page.insert_text((x0 + 4, y0 + 12.5), cell, fontsize=9.5, fontname="helv")
        self.y = top + height * len(rows) + 22


def build_pdf(content: dict[str, Any]) -> bytes:
    w = _PdfWriter()
    w.line(content["name"], size=16, bold=True)
    for line in content.get("contact", []):
        w.line(line, size=9.5)
    for link in content.get("links", []):
        w.ensure(_LINE)
        y = w.y
        w.line(link["text"], size=9.5, color=(0, 0, 0.8))
        rect = pymupdf.Rect(_LEFT, y - 10, _LEFT + 6 * len(link["text"]), y + 3)
        w.page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": link["url"]})
    if content.get("hidden_text"):
        w.line(content["hidden_text"], size=6, color=(1, 1, 1))
    for section in content["sections"]:
        w.gap(10)
        if section.get("heading"):
            w.ensure(_LINE * 3)
            w.line(section["heading"].upper(), size=12, bold=True)
        for paragraph in section.get("paragraphs", []):
            w.wrapped(paragraph)
            w.gap(4)
        for bullet in section.get("bullets", []):
            w.wrapped(bullet, bullet=True)
        if section.get("table"):
            w.gap(6)
            w.table(section["table"])
    if content.get("image_only_page"):
        page = w.doc.new_page()
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 120, 160), 0)
        pix.clear_with(200)
        page.insert_image(pymupdf.Rect(60, 60, 540, 700), pixmap=pix)
    data = w.doc.tobytes(garbage=3, deflate=True)
    w.doc.close()
    return data


def build_pdf_pages(page_texts: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in page_texts:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def build_encrypted_pdf(text: str = "Secret resume") -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text)
    buffer = io.BytesIO()
    doc.save(buffer, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    doc.close()
    return buffer.getvalue()


def build_image_only_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), 0)
    pix.clear_with(180)
    page.insert_image(pymupdf.Rect(72, 72, 300, 300), pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data


# ------------------------------------------------------------------ DOCX


def add_hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bookmark_link(paragraph, anchor: str, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_docx(content: dict[str, Any], use_heading_styles: bool = True) -> bytes:
    document = docx.Document()
    document.add_paragraph(content["name"], style="Title" if use_heading_styles else None)
    for line in content.get("contact", []):
        document.add_paragraph(line)
    for link in content.get("links", []):
        add_hyperlink(document.add_paragraph(), link["url"], link["text"])
    for section in content["sections"]:
        if section.get("heading"):
            if use_heading_styles:
                document.add_heading(section["heading"], level=1)
            else:
                document.add_paragraph(section["heading"].upper())
        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets", []):
            document.add_paragraph(bullet, style="List Bullet")
        if section.get("table"):
            rows = section["table"]
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            for r_index, row in enumerate(rows):
                for c_index, cell in enumerate(row):
                    table.cell(r_index, c_index).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_docx_paragraphs(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
